from docx import Document
from docx.oxml.ns import qn
import re

d = Document(r"D:\phongthi247\scripts\_toan_src.docx")
ps = d.paragraphs

print("TABLES:", len(d.tables))
xml = d.element.xml
print("sdt:", xml.count("<w:sdt>"), "tbl:", xml.count("<w:tbl>"), "txbx:", xml.count("txbxContent"))

# Walk body children in order, print type + text for the P2 region
body = d.element.body
order = []
for child in body.iterchildren():
    tag = child.tag.split('}')[-1]
    if tag == 'p':
        txt = "".join(t.text or "" for t in child.iter(qn('w:t')))
        pPr = child.find(qn('w:pPr'))
        num = pPr is not None and pPr.find(qn('w:numPr')) is not None
        order.append(('p', num, txt))
    elif tag == 'tbl':
        rows = []
        for tr in child.findall(qn('w:tr')):
            cells = []
            for tc in tr.findall(qn('w:tc')):
                cells.append("".join(t.text or "" for t in tc.iter(qn('w:t'))))
            rows.append(" | ".join(cells))
        order.append(('tbl', False, " ;; ".join(rows)))
    else:
        order.append((tag, False, ''))

start = next(i for i, o in enumerate(order) if o[2].strip().startswith('PHẦN II'))
end = next(i for i, o in enumerate(order) if o[2].strip().startswith('PHẦN III'))
print(f"P2 range: {start}..{end}")
for i in range(start, min(end + 3, len(order))):
    kind, num, txt = order[i]
    print(f"{i} [{kind}{'/num' if num else ''}] {txt[:150]}")
