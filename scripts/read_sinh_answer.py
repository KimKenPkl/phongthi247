from docx import Document
from pathlib import Path
import glob

down = Path(r"C:\Users\LENOVO\Downloads")
target = None
for f in down.iterdir():
    if f.suffix.lower() == ".docx" and "CUONG GHKI" in f.name:
        target = f
        print("TARGET FILE:", repr(f.name))
        break

d = Document(target)
full = []
for para in d.paragraphs:
    t = para.text.strip()
    if t: full.append(("P", t))
for ti, table in enumerate(d.tables):
    for row in table.rows:
        txt = " | ".join(c.text.strip() for c in row.cells)
        if txt: full.append((f"T{ti}", txt))

content = "\n".join(t for _, t in full)
print("TOTAL:", len(content))
import re
for tag, t in full:
    if re.search(r"Câu\s*1:|Câu\s*12:|Câu\s*6\.|ĐÚNG SAI|TRẢ LỜI NGẮN", t):
        print(f"[{tag}] {t[:250]}")
        print("-"*70)
