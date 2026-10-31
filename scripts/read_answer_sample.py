from docx import Document
from pathlib import Path

f = Path(r"C:\Users\LENOVO\Downloads\DAP AN D? CUONG GHKI_K12_26-27.docx")
d = Document(f)

full = []
for para in d.paragraphs:
    t = para.text.strip()
    if t: full.append(("P", t))
for ti, table in enumerate(d.tables):
    for row in table.rows:
        txt = " | ".join(c.text.strip() for c in row.cells)
        if txt: full.append((f"T{ti}", txt))

content = "\n".join(t for _, t in full)
print("TOTAL:", len(content))

# Find relevant questions: D/S Q1, Q12; P3 Q6
import re
for tag, t in full:
    if re.search(r"Câu\s*1:|Câu\s*12:|Câu\s*6\.|ĐÚNG SAI|TRẢ LỜI NGẮN", t):
        print(f"[{tag}] {t[:200]}")
        print()
