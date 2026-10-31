from docx import Document
from pathlib import Path

f = Path(r"C:\Users\LENOVO\Downloads\Đề cương ôn tập GK 1-Đáp án (2).docx")
d = Document(f)

full = []
for para in d.paragraphs:
    t = para.text.strip()
    if t: full.append(("P", t))
for ti, table in enumerate(d.tables):
    for row in table.rows:
        txt = " | ".join(c.text.strip() for c in row.cells)
        if txt: full.append((f"T{ti}", txt))

content = "\n".join(t for _, t in full)
print("TOTAL:", len(content))

for tag, t in full:
    if any(k in t for k in ["Cau", "Câu", "ĐÚNG", "TRẢ LỜI", "PHẦN", "Câu 1", "Câu 6", "6."]):
        print(f"[{tag}] {t[:180]}")
